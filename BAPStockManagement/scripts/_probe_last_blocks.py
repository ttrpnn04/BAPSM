# -*- coding: utf-8 -*-
"""Quick check: stock totals around last blocks for A025 16\"."""
from pathlib import Path
import openpyxl

path = next(p for p in Path(r"C:\Users\Admin\Downloads").glob("*0769*.xlsx") if not p.name.startswith("~$"))
wb = openpyxl.load_workbook(path, data_only=True)
ws = wb.active
row = 5
# last two stock starts
for start in (420, 452):
    vals = []
    for off in range(14):
        h = str(ws.cell(2, start + off).value or "").strip()
        v = ws.cell(row, start + off).value or 0
        try:
            v = int(v)
        except Exception:
            v = 0
        if v:
            vals.append(f"{h}={v}")
    total = ws.cell(row, start + 15).value
    print(f"start={start} {' '.join(vals)} total_col={total}")

import datetime as dt
from pathlib import Path

import openpyxl
import pyodbc


EXCEL_PATH = Path(r"C:\Users\pigeo\OneDrive\Desktop\สต็อกจักรยาน 0669.xlsx")
DB_CONN_CANDIDATES = [
    (
        "Driver={ODBC Driver 17 for SQL Server};"
        r"Server=localhost\SQLEXPRESS;"
        "Database=BAPStockDB;"
        "Trusted_Connection=yes;"
        "TrustServerCertificate=yes;"
    ),
    (
        "Driver={ODBC Driver 17 for SQL Server};"
        r"Server=localhost\SQLEXPRESS03;"
        "Database=BAPStockDB;"
        "Trusted_Connection=yes;"
        "TrustServerCertificate=yes;"
    ),
]
IMPORT_REF = "EXCEL-INITIAL-0669"
CATEGORY_NAME = "กลุ่มจักรบยาน-1(LION)"
CREATED_BY = None


def parse_latest_stock_rows():
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws = wb[wb.sheetnames[0]]

    stock_starts = []
    for c in range(1, ws.max_column + 1):
        v = ws.cell(1, c).value
        if isinstance(v, str) and "สต็อก" in v:
            stock_starts.append(c)

    blocks = []
    for start in stock_starts:
        colors = []
        c = start
        while c <= ws.max_column:
            header = ws.cell(2, c).value
            if header is None:
                break
            hs = str(header).strip()
            if hs in ("ลัง", "คงเหลือ", "รวม", "ยอดรวม"):
                break
            if "บิล" in hs or hs == "เข้า":
                break
            colors.append((c, hs))
            c += 1
        if len(colors) >= 3:
            blocks.append((start, colors))

    if not blocks:
        raise RuntimeError("ไม่พบ block สต็อกในไฟล์ Excel")

    _, colors = blocks[-1]

    rows = []
    for r in range(3, ws.max_row + 1):
        sku = ws.cell(r, 2).value
        product_name = ws.cell(r, 3).value
        if not sku or not product_name:
            continue
        sku = str(sku).strip()
        product_name = str(product_name).strip()
        if not sku or not product_name:
            continue

        variants = []
        for col, color_name in colors:
            raw = ws.cell(r, col).value
            try:
                qty = int(float(raw or 0))
            except Exception:
                qty = 0
            if qty > 0:
                variants.append((color_name, qty))

        if variants:
            rows.append((sku, product_name, variants))

    return rows


def get_or_create_category(cur):
    cur.execute(
        "SELECT CategoryID FROM Categories WHERE CategoryName = ?",
        CATEGORY_NAME,
    )
    row = cur.fetchone()
    if row:
        return row[0]

    cur.execute(
        """
        INSERT INTO Categories (CategoryName, VariantLabel, SortOrder, IsActive)
        VALUES (?, N'สี', 999, 1)
        """,
        CATEGORY_NAME,
    )
    cur.execute(
        "SELECT CategoryID FROM Categories WHERE CategoryName = ?",
        CATEGORY_NAME,
    )
    return int(cur.fetchone()[0])


def get_receive_txn_type(cur):
    cur.execute(
        """
        SELECT TOP 1 TransactionTypeID
        FROM TransactionTypes
        WHERE Direction = 1 AND IsActive = 1
        ORDER BY TransactionTypeID
        """
    )
    row = cur.fetchone()
    if row:
        return row[0]

    cur.execute(
        """
        INSERT INTO TransactionTypes (TypeName, Direction, IsActive)
        VALUES (N'รับเข้า', 1, 1)
        """
    )
    cur.execute(
        "SELECT TOP 1 TransactionTypeID FROM TransactionTypes WHERE TypeName = N'รับเข้า' ORDER BY TransactionTypeID DESC"
    )
    return int(cur.fetchone()[0])


def get_or_create_product(cur, category_id, sku, product_name):
    cur.execute(
        """
        SELECT ProductID
        FROM Products
        WHERE SKU = ? AND ProductName = ?
        """,
        sku,
        product_name,
    )
    row = cur.fetchone()
    if row:
        product_id = int(row[0])
        cur.execute(
            """
            UPDATE Products
            SET CategoryID = ?, IsActive = 1, UpdatedAt = SYSDATETIME()
            WHERE ProductID = ?
            """,
            category_id,
            product_id,
        )
        return product_id

    cur.execute(
        """
        INSERT INTO Products (CategoryID, SKU, ProductName, Unit, IsActive, Note)
        VALUES (?, ?, ?, N'คัน', 1, N'Imported from Excel 0669')
        """,
        category_id,
        sku,
        product_name,
    )
    cur.execute(
        "SELECT ProductID FROM Products WHERE SKU = ? AND ProductName = ?",
        sku,
        product_name,
    )
    return int(cur.fetchone()[0])


def get_or_create_variant(cur, product_id, variant_name):
    cur.execute(
        """
        SELECT VariantID
        FROM ProductVariants
        WHERE ProductID = ? AND VariantName = ?
        """,
        product_id,
        variant_name,
    )
    row = cur.fetchone()
    if row:
        variant_id = int(row[0])
        cur.execute(
            "UPDATE ProductVariants SET IsActive = 1 WHERE VariantID = ?",
            variant_id,
        )
        return variant_id

    cur.execute(
        "SELECT ISNULL(MAX(SortOrder), 0) + 1 FROM ProductVariants WHERE ProductID = ?",
        product_id,
    )
    next_sort = int(cur.fetchone()[0])

    cur.execute(
        """
        INSERT INTO ProductVariants (ProductID, VariantName, SortOrder, IsActive)
        VALUES (?, ?, ?, 1)
        """,
        product_id,
        variant_name,
        next_sort,
    )
    cur.execute(
        "SELECT VariantID FROM ProductVariants WHERE ProductID = ? AND VariantName = ?",
        product_id,
        variant_name,
    )
    variant_id = int(cur.fetchone()[0])

    cur.execute(
        """
        INSERT INTO StockBalances (VariantID, QtyPieces, QtyCases)
        VALUES (?, 0, 0)
        """,
        variant_id,
    )

    return variant_id


def main():
    if not EXCEL_PATH.exists():
        raise FileNotFoundError(f"ไม่พบไฟล์ {EXCEL_PATH}")

    rows = parse_latest_stock_rows()
    if not rows:
        raise RuntimeError("ไม่พบข้อมูลสต็อกที่มากกว่า 0 ในไฟล์")

    last_error = None
    conn = None
    for conn_str in DB_CONN_CANDIDATES:
        try:
            conn = pyodbc.connect(conn_str, timeout=5)
            break
        except Exception as ex:
            last_error = ex
    if conn is None:
        raise RuntimeError(f"เชื่อม SQL Server ไม่ได้: {last_error}")
    conn.autocommit = False
    cur = conn.cursor()

    try:
        # Prevent accidental duplicate import
        cur.execute(
            "SELECT COUNT(*) FROM StockTransactions WHERE RefNo = ?",
            IMPORT_REF,
        )
        if int(cur.fetchone()[0]) > 0:
            raise RuntimeError(f"พบการนำเข้าซ้ำแล้ว (RefNo={IMPORT_REF})")

        category_id = get_or_create_category(cur)
        txn_type_id = get_receive_txn_type(cur)

        tx_date = dt.date.today()
        products_count = 0
        variants_count = 0
        tx_count = 0
        total_pieces = 0

        for sku, product_name, variants in rows:
            product_id = get_or_create_product(cur, category_id, sku, product_name)
            products_count += 1

            for variant_name, qty in variants:
                variant_id = get_or_create_variant(cur, product_id, variant_name)
                variants_count += 1
                total_pieces += qty

                cur.execute(
                    """
                    INSERT INTO StockTransactions
                    (VariantID, TransactionTypeID, TxnDate, QtyPieces, QtyCases, RefNo, Note, CreatedBy)
                    VALUES (?, ?, ?, ?, 0, ?, N'Initial import from Excel: สต็อกจักรยาน 0669.xlsx', ?)
                    """,
                    variant_id,
                    txn_type_id,
                    tx_date,
                    qty,
                    IMPORT_REF,
                    CREATED_BY,
                )
                tx_count += 1

        conn.commit()
        print(
            f"IMPORTED OK | products_rows={products_count} | variants_with_stock={variants_count} "
            f"| transactions={tx_count} | total_pieces={total_pieces} | ref={IMPORT_REF}"
        )
    except Exception:
        conn.rollback()
        raise
    finally:
        cur.close()
        conn.close()


if __name__ == "__main__":
    main()

# -*- coding: utf-8 -*-
"""Find rightmost stock block values for BLB-A025 16\" in full 0769 file."""
from pathlib import Path
import openpyxl

path = next(p for p in Path(r"C:\Users\Admin\Downloads").glob("*0769*.xlsx") if not p.name.startswith("~$"))
wb = openpyxl.load_workbook(path, data_only=True)
ws = wb[wb.sheetnames[0]]

# find row
row = None
for r in range(1, ws.max_row + 1):
    sku = ws.cell(r, 2).value
    name = ws.cell(r, 3).value
    if sku and "BLB-A025" in str(sku) and name and str(name).startswith("16"):
        row = r
        break

# find all stock block starts (row1 contains สต็อก, row2 = แดง)
starts = []
for c in range(1, ws.max_column + 1):
    t = str(ws.cell(1, c).value or "").strip()
    h = str(ws.cell(2, c).value or "").strip()
    if "สต็อก" in t and h.startswith("แดง"):
        starts.append(c)

out = []
out.append(f"file={path.name} row={row} stock_starts={starts}")

for start in starts:
    # read 14 color cols + ลัง + maybe total after
    colors = {}
    for off, expected in enumerate(["แดง","น้ำเงิน/ชม","ฟ้า","ส้ม","เขียว","เขียวอ่อน","ชมพูอ่อน","ชมเข้ม/ชม/ชม","เหลือง/ทอง/ครีม","ม่วง","วัว/ขาว /แดง","ดำ/ธงฟ้า","ตาล","เทา","ลัง"]):
        c = start + off
        h = str(ws.cell(2, c).value or "").strip()
        v = ws.cell(row, c).value
        colors[h or expected] = v
    total_col = start + 15  # often empty header total
    total_v = ws.cell(row, total_col).value
    red = colors.get("แดง") or colors.get(list(colors.keys())[0])
    out.append(f"start={start} แดง={colors.get('แดง')} ฟ้า={next((v for k,v in colors.items() if 'ฟ้า' in k and 'ชม' not in k and 'ธง' not in k), None)} เขียว={next((v for k,v in colors.items() if k.startswith('เขียว') and 'อ่อน' not in k), None)} ชมเข้ม={next((v for k,v in colors.items() if 'ชมเข้ม' in k), None)} ม่วง={next((v for k,v in colors.items() if 'ม่วง' in k), None)} total_next={total_v}")

# Also simulate OLD FindQuantityBlocks last block
# and NEW logic

def find_blocks_old():
    blocks = []
    for col in range(1, ws.max_column + 1):
        title = str(ws.cell(1, col).value or "").strip()
        first = str(ws.cell(2, col).value or "").strip()
        if "สต็อก" not in title or not first:
            continue
        cols = []
        for q in range(col, ws.max_column + 1):
            h = str(ws.cell(2, q).value or "").strip()
            if not h:
                break
            cols.append((q, h))
        if len(cols) >= 2:
            blocks.append((col, cols))
    return blocks

def find_blocks_new():
    blocks = []
    prev = False
    for col in range(1, ws.max_column + 1):
        title = str(ws.cell(1, col).value or "").strip()
        matched = "สต็อก" in title
        is_start = matched and not prev
        prev = matched
        if not is_start:
            continue
        first = str(ws.cell(2, col).value or "").strip()
        if not first:
            continue
        cols = []
        for q in range(col, ws.max_column + 1):
            section = str(ws.cell(1, q).value or "").strip()
            if q > col and section and "สต็อก" not in section:
                break
            h = str(ws.cell(2, q).value or "").strip()
            if not h:
                break
            if h in ("คงเหลือ", "รวม", "ยอดรวม", "รวมจำนวน", "รวมทั้งหมด", "total"):
                continue
            cols.append((q, h))
        if len(cols) >= 2:
            blocks.append((col, cols))
    return blocks

old_b = find_blocks_old()
new_b = find_blocks_new()
out.append(f"OLD blocks={len(old_b)} last_start={old_b[-1][0] if old_b else None} last_ncols={len(old_b[-1][1]) if old_b else None} last_headers={[h for _,h in old_b[-1][1][:5]] if old_b else None}")
out.append(f"NEW blocks={len(new_b)} last_start={new_b[-1][0] if new_b else None} last_ncols={len(new_b[-1][1]) if new_b else None} last_headers={[h for _,h in new_b[-1][1][:5]] if new_b else None}")

for label, blocks in [("OLD", old_b), ("NEW", new_b)]:
    if not blocks:
        continue
    # candidates ending with ลัง
    cands = [b for b in blocks if b[1][-1][1].strip() == "ลัง"]
    if not cands:
        out.append(f"{label} no ลัง candidates")
        continue
    # rightmost
    col, cols = max(cands, key=lambda b: b[1][0][0])
    out.append(f"{label} selected start={col} headers={[h for _,h in cols]}")
    for q, h in cols:
        if h.strip() in ("แดง", "ฟ้า", "เขียว", "ชมเข้ม/ชม/ชม", "ม่วง") or "ฟ้า" in h or "เขียว" in h.replace("อ่อน","") and "อ่อน" not in h:
            v = ws.cell(row, q).value
            out.append(f"  {h!r}={v}")

Path(r"C:\Users\Admin\Documents\GitHub\BAPSM\BAPStockManagement\scripts\_excel_probe2.txt").write_text("\n".join(out), encoding="utf-8")
print("done", len(out))

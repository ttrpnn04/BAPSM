from pathlib import Path

import openpyxl
import pyodbc

EXCEL_PATH = Path(r"C:\Users\pigeo\OneDrive\Desktop\สต็อกจักรยาน 0669.xlsx")
DB_CONN = (
    "Driver={ODBC Driver 17 for SQL Server};"
    r"Server=localhost\SQLEXPRESS03;"
    "Database=BAPStockDB;"
    "Trusted_Connection=yes;"
    "TrustServerCertificate=yes;"
)
CATEGORY_NAME = "กลุ่มจักรบยาน-1(LION)"


def parse_latest_stock_columns(ws):
    stock_starts = []
    for c in range(1, ws.max_column + 1):
        value = ws.cell(1, c).value
        if isinstance(value, str) and "สต็อก" in value:
            stock_starts.append(c)

    blocks = []
    for start in stock_starts:
        cols = []
        col = start
        while col <= ws.max_column:
            header = ws.cell(2, col).value
            if header is None:
                break
            hs = str(header).strip()
            if hs in ("ลัง", "คงเหลือ", "รวม", "ยอดรวม") or "บิล" in hs or hs == "เข้า":
                break
            cols.append(col)
            col += 1
        if len(cols) >= 3:
            blocks.append(cols)

    if not blocks:
        raise RuntimeError("ไม่พบ block สต็อก")
    return blocks[-1]


def main():
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws = wb[wb.sheetnames[0]]
    stock_cols = parse_latest_stock_columns(ws)

    conn = pyodbc.connect(DB_CONN)
    conn.autocommit = False
    cur = conn.cursor()

    cur.execute("SELECT CategoryID FROM Categories WHERE CategoryName = ?", CATEGORY_NAME)
    row = cur.fetchone()
    if row:
        category_id = int(row[0])
    else:
        cur.execute(
            """
            INSERT INTO Categories (CategoryName, VariantLabel, SortOrder, IsActive)
            VALUES (?, N'สี', 999, 1)
            """,
            CATEGORY_NAME,
        )
        cur.execute("SELECT CategoryID FROM Categories WHERE CategoryName = ?", CATEGORY_NAME)
        category_id = int(cur.fetchone()[0])

    added = 0
    updated = 0
    skipped_has_stock = 0

    for r in range(3, ws.max_row + 1):
        sku = ws.cell(r, 2).value
        product_name = ws.cell(r, 3).value
        if not sku or not product_name:
            continue

        sku = str(sku).strip()
        product_name = str(product_name).strip()
        if not sku or not product_name:
            continue

        has_stock = False
        for c in stock_cols:
            v = ws.cell(r, c).value
            try:
                qty = int(float(v or 0))
            except Exception:
                qty = 0
            if qty > 0:
                has_stock = True
                break

        if has_stock:
            skipped_has_stock += 1
            continue

        cur.execute(
            "SELECT ProductID FROM Products WHERE SKU = ? AND ProductName = ?",
            sku[:50],
            product_name[:300],
        )
        existing = cur.fetchone()
        if existing:
            updated += 1
            cur.execute(
                """
                UPDATE Products
                SET CategoryID = ?, IsActive = 1, UpdatedAt = SYSDATETIME()
                WHERE ProductID = ?
                """,
                category_id,
                int(existing[0]),
            )
        else:
            added += 1
            cur.execute(
                """
                INSERT INTO Products (CategoryID, SKU, ProductName, Unit, IsActive, Note)
                VALUES (?, ?, ?, N'คัน', 1, N'Imported zero-stock master from Excel 0669')
                """,
                category_id,
                sku[:50],
                product_name[:300],
            )

    conn.commit()
    conn.close()
    print(f"ZERO_STOCK_MASTER_DONE | added={added} | updated={updated} | skipped_has_stock={skipped_has_stock}")


if __name__ == "__main__":
    main()

from pathlib import Path

import openpyxl
import pyodbc

EXCEL_PATH = Path(r"C:\Users\pigeo\OneDrive\Desktop\สต็อกจักรยาน 0669.xlsx")
DB_CONN = (
    "Driver={ODBC Driver 17 for SQL Server};"
    r"Server=localhost\SQLEXPRESS03;"
    "Database=BAPStockDB;"
    "Trusted_Connection=yes;"
    "TrustServerCertificate=yes;"
)


def main():
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws = wb[wb.sheetnames[0]]

    excel_rows = set()
    for r in range(3, ws.max_row + 1):
        sku = ws.cell(r, 2).value
        name = ws.cell(r, 3).value
        if not sku or not name:
            continue
        sku = str(sku).strip()
        name = str(name).strip()
        if not sku or not name:
            continue
        excel_rows.add((sku[:50], name[:300]))

    conn = pyodbc.connect(DB_CONN)
    cur = conn.cursor()
    db_matched = 0
    for sku, name in excel_rows:
        cur.execute(
            "SELECT COUNT(1) FROM Products WHERE SKU = ? AND ProductName = ?",
            sku,
            name,
        )
        if int(cur.fetchone()[0]) > 0:
            db_matched += 1

    cur.execute(
        """
        SELECT c.CategoryName, COUNT(1)
        FROM Products p
        INNER JOIN Categories c ON c.CategoryID = p.CategoryID
        WHERE p.SKU LIKE 'BL%'
        GROUP BY c.CategoryName
        ORDER BY c.CategoryName
        """
    )
    groups = cur.fetchall()
    conn.close()

    print(f"excel_unique={len(excel_rows)} db_matched={db_matched}")
    for name, count in groups:
        safe_name = str(name).encode("unicode_escape").decode("ascii")
        print(f"{safe_name}: {count}")


if __name__ == "__main__":
    main()
